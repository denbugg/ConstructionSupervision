"""Справочник работ через API: загрузка XLSX, восстановленные коды, фильтры, порядок."""

import io
from datetime import datetime

from openpyxl import Workbook

WORK_TYPES = "/api/v1/plan/work-types"
V = "˅"


def _xlsx(rows: list[list]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["Справочник Видов Работ"])
    sheet.append(["№ п/п", "Вид работ", "Жильё", "Дороги"])
    for row in rows:
        sheet.append(row)
        if isinstance(row[0], datetime):
            sheet.cell(sheet.max_row, 1).number_format = "d\\.m\\."
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


REFERENCE = _xlsx(
    [
        ["10.", "Подготовка территории", V, V],
        [datetime(2025, 2, 10), "Вынос сетей", V, V],
        [None, "Вынос сетей: теплосеть", V, None],
        [datetime(2025, 10, 10), "Геодезические знаки", V, V],
        [12, "Строительно-монтажные работы", V, V],
    ]
)


async def _import(client, content=REFERENCE, name="spr.xlsx"):
    return await client.post(
        f"{WORK_TYPES}/import",
        files={"file": (name, content, "application/octet-stream")},
    )


async def test_загрузка_справочника_и_восстановленные_коды(client):
    response = await _import(client)

    assert response.status_code == 200, response.text
    assert response.json() == {
        "work_types": 5,
        "rows_without_code": 1,
        "restored_codes": [
            {"row": 4, "cell": "10.02.2025", "code": "10.2"},
            {"row": 6, "cell": "10.10.2025", "code": "10.10"},
        ],
    }
    page = (await client.get(WORK_TYPES)).json()
    assert page["total"] == 5
    assert [w["code"] for w in page["items"]] == ["10", "10.2", "10.2/1", "10.10", "12"]
    heat = page["items"][2]
    assert heat == {
        "code": "10.2/1",
        "name": "Вынос сетей: теплосеть",
        "level": 4,
        "parent_code": "10.2",
        "applicable": {"Жильё": True, "Дороги": False},
        "source": "spr.xlsx, лист «Sheet», строка 5",
    }


async def test_фильтры_уровня_и_родителя(client):
    await _import(client)

    level_2 = (await client.get(WORK_TYPES, params={"level": 2})).json()
    children = (await client.get(WORK_TYPES, params={"parent_code": "10.2"})).json()

    assert [w["code"] for w in level_2["items"]] == ["10.2", "10.10"]
    assert [w["code"] for w in children["items"]] == ["10.2/1"]
    assert (await client.get(WORK_TYPES, params={"level": 5})).status_code == 422


async def test_повторная_загрузка_заменяет_справочник(client):
    await _import(client)

    await _import(client, _xlsx([["12", "Строительно-монтажные работы", V, V]]))

    assert [w["code"] for w in (await client.get(WORK_TYPES)).json()["items"]] == ["12"]


async def test_ошибка_называет_строку_и_ничего_не_меняет(client):
    await _import(client)
    bad = _xlsx([["10.", "Подготовка", "v", V]])

    response = await _import(client, bad)

    assert response.status_code == 400
    error = response.json()["error"]
    assert error["code"] == "WORK_TYPES_IMPORT_INVALID"
    assert error["details"]["errors"][0] | {"message": None} == {
        "row": 3,
        "column": "Жильё",
        "message": None,
    }
    assert (await client.get(WORK_TYPES)).json()["total"] == 5


async def test_не_xlsx(client):
    response = await _import(client, b"%PDF", name="spr.pdf")

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "WORK_TYPES_IMPORT_INVALID"
