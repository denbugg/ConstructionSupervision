"""Парсер справочника работ: испорченные коды, строки без кода, отметки, ошибки по строкам."""

import io
from datetime import datetime

import pytest
from openpyxl import Workbook
from src.core.reference_import import (
    MARK,
    ReferenceImportError,
    code_sort_key,
    parse_work_types,
    read_sheet,
)

from tests.conftest import SERVICE_ROOT

REFERENCE = SERVICE_ROOT.parents[1] / "data/reference/Сводный перечень строительных работ_ЛТЦ.xlsx"
HEADER = ["№ п/п", "Вид работ", "Жильё", "Дороги"]
V = MARK


def _xlsx(rows: list[list]) -> bytes:
    """Книга как у заказчика: заголовок справочника, шапка в строке 3, данные с 4-й."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист1"
    sheet.append([])
    sheet.append(["Справочник Видов Работ"])
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
        code = sheet.cell(sheet.max_row, 1)
        if isinstance(code.value, datetime):
            code.number_format = "d\\.m\\."  # так столбец кодов отформатирован в файле
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _parse(rows: list[list]):
    sheet, table = read_sheet(_xlsx(rows))
    return parse_work_types(table, filename="spr.xlsx", sheet=sheet)


def _issues(rows: list[list]) -> list[tuple[int | None, str | None]]:
    with pytest.raises(ReferenceImportError) as error:
        _parse(rows)
    return [(i.row, i.column) for i in error.value.issues]


def test_коды_из_дат_числа_и_с_точкой_восстанавливаются():
    items = _parse(
        [
            ["10.", "Подготовка территории", V, V],
            [datetime(2025, 1, 10), "Отселение домов", V, V],
            [datetime(2025, 11, 10), "Обустройство площадки", V, V],
            ["10.11.1.", "Ограждение площадки", None, V],
            [12, "Строительно-монтажные работы", V, V],
            [datetime(2025, 3, 12), "Подземная часть", V, None],
            ["12.3.1.", "Котлован", V, V],
        ]
    )

    assert [(i.code, i.level, i.parent_code) for i in items] == [
        ("10", 1, None),
        ("10.1", 2, "10"),
        ("10.11", 2, "10"),
        ("10.11.1", 3, "10.11"),
        ("12", 1, None),
        ("12.3", 2, "12"),
        ("12.3.1", 3, "12.3"),
    ]
    assert [i.restored_from for i in items if i.restored_from] == [
        "10.01.2025",
        "10.11.2025",
        "12.03.2025",
    ]
    assert all(i.code_in_file for i in items)


def test_строки_без_кода_4_го_уровня_под_ближайшим_кодом_сверху():
    items = _parse(
        [
            ["10.", "Подготовка территории", V, V],
            [datetime(2025, 2, 10), "Вынос сетей", V, V],
            [None, "Вынос сетей: теплосеть", V, V],
            [None, "Вынос сетей: водоснабжение", V, V],
            [datetime(2025, 3, 10), "Снос домов", V, V],
            ["10.3.1.", "Снос жилых", V, V],
            [None, "Демонтаж перекрытий", V, None],
        ]
    )

    without_code = [(i.code, i.level, i.parent_code) for i in items if not i.code_in_file]
    assert without_code == [
        ("10.2/1", 4, "10.2"),
        ("10.2/2", 4, "10.2"),
        ("10.3.1/1", 4, "10.3.1"),
    ]


def test_отметки_по_столбцам_шапки_и_источник_строки():
    item = _parse([["10.", "Подготовка территории", V, None]])[0]

    assert item.applicable == {"Жильё": True, "Дороги": False}
    assert item.source == "spr.xlsx, лист «Лист1», строка 4"


def test_перенос_строки_в_наименовании_схлопывается():
    item = _parse([["10.", "Монтаж панелей \n (нащельники)", V, V]])[0]

    assert item.name == "Монтаж панелей (нащельники)"


def test_ошибки_собираются_по_всему_файлу_с_номерами_строк():
    issues = _issues(
        [
            [None, "Строка до первого кода", V, V],  # 4
            ["10.", "", V, V],  # 5: нет наименования
            ["10.1.", "Отселение", "v", V],  # 6: латинская v вместо отметки
            [12.1, "Код-дробь", V, V],  # 7: «12.1» или «12.10» — не угадать
            ["abc", "Мусор в коде", V, V],  # 8
            ["10.1.", "Повтор кода", V, V],  # 9
            ["11.2.", "Родителя 11 нет", V, V],  # 10
            ["10.1.1.1.1.", "Пятый уровень", V, V],  # 11
        ]
    )

    assert issues == [
        (4, "№ п/п"),
        (5, "Вид работ"),
        (6, "Жильё"),
        (7, "№ п/п"),
        (8, "№ п/п"),
        (9, "№ п/п"),
        (10, "№ п/п"),
        (11, "№ п/п"),
        (11, "№ п/п"),
    ]


def test_шапка_без_столбца_кодов_или_типов_объектов():
    book = Workbook()
    book.active.append(["Вид работ", "Жильё"])
    out = io.BytesIO()
    book.save(out)
    sheet, table = read_sheet(out.getvalue())

    with pytest.raises(ReferenceImportError, match="№ п/п"):
        parse_work_types(table, filename="x.xlsx", sheet=sheet)
    with pytest.raises(ReferenceImportError, match="типов объектов"):
        parse_work_types([(1, ["№ п/п", "Вид работ"])], filename="x.xlsx", sheet="Лист1")


def test_не_xlsx_и_пустой_справочник():
    with pytest.raises(ReferenceImportError, match="не читается"):
        read_sheet(b"%PDF")
    with pytest.raises(ReferenceImportError, match="ни одного"):
        _parse([])


def test_порядок_кодов_как_в_файле_а_не_строковый():
    codes = ["10.10", "10.2.1", "10", "10.2/2", "10.2", "12", "10.2/10", "10.2/1", "10.11.1"]

    assert sorted(codes, key=code_sort_key) == [
        "10",
        "10.2",
        "10.2/1",
        "10.2/2",
        "10.2/10",
        "10.2.1",
        "10.10",
        "10.11.1",
        "12",
    ]


@pytest.mark.skipif(not REFERENCE.exists(), reason="нет data/reference: тест из рабочей копии")
def test_настоящий_справочник_заказчика():
    """Числа — из data/README.md: 377 строк данных, 250 без кода, 19 кодов-дат."""
    sheet, table = read_sheet(REFERENCE.read_bytes())

    items = parse_work_types(table, filename=REFERENCE.name, sheet=sheet)

    assert len(items) == 377
    assert sum(not i.code_in_file for i in items) == 250
    restored = [i.code for i in items if i.restored_from]
    assert restored == [f"10.{n}" for n in range(1, 13)] + [f"12.{n}" for n in range(1, 8)]
    by_code = {i.code: i for i in items}
    assert (by_code["12"].level, by_code["12.3.1"].name) == (1, "Устройство котлована")
    assert by_code["12.6.13.1"].level == 4
    assert all(len(i.applicable) == 9 for i in items)
    assert by_code["10.11.1"].applicable["Дороги"] and not by_code["10.11.1"].applicable["Жильё"]
    assert [i.code for i in sorted(items, key=lambda i: code_sort_key(i.code))] == [
        i.code for i in items
    ]
