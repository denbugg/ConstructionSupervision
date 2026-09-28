"""Импорт календарного графика из CSV или XLSX (ТЗ, п. 3: код, наименование, начало, окончание,
необязательно тип участка, визуальная стадия, связи).

Строка с кодом из шаблона этапов получает оттуда недостающее: фазу, тип участка, визуальную
стадию, связи и правило «этап → техника». Столбцы файла главнее шаблона. Ошибки собираются
по всем строкам сразу, с номером строки файла: исправлять файл по одной ошибке за загрузку
долго.
"""

import csv
import io
import re
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from typing import Any
from uuid import NAMESPACE_URL, uuid5

from src.core.calendar import WorkCalendar, count_working_days
from src.core.cpm import CpmError, CpmStage, Link, order_stages
from src.core.stages import WORK_ROLE
from src.core.templates import TemplateStage, Vocabulary

# Название поля → допустимые заголовки столбца (без учёта регистра).
COLUMNS = {
    "code": ("код", "code"),
    "name": ("наименование", "name"),
    "start": ("начало", "start", "plan_start"),
    "end": ("окончание", "end", "plan_end"),
    "zone_type": ("тип участка", "zone_type"),
    "visual_stage": ("визуальная стадия", "visual_stage"),
    "predecessors": ("связи", "predecessors"),
    "phase": ("фаза", "phase"),
}
REQUIRED = ("code", "name", "start", "end")
# Связь: «12.3.1», «12.3.1 FS», «12.3.1 SS+2». Несколько — через «;» или «,».
LINK_PATTERN = re.compile(r"^(?P<code>\S+)(?:\s+(?P<type>FS|SS|FF|SF)(?P<lag>[+-]\d+)?)?$", re.I)
NO_LINKS = ("-", "—", "нет")
DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y")


@dataclass(frozen=True)
class ImportIssue:
    row: int | None
    column: str | None
    message: str


class PlanImportError(ValueError):
    """Файл графика не принят: список проблем с номерами строк."""

    def __init__(self, issues: Sequence[ImportIssue]) -> None:
        self.issues = list(issues)
        first = self.issues[0]
        where = f"строка {first.row}: " if first.row else ""
        super().__init__(f"{where}{first.message}")


@dataclass(frozen=True)
class ImportedLink:
    code: str
    type: str
    lag_days: int
    # Связь из шаблона, а не из файла: на этап, которого в файле нет, она просто не ставится.
    from_template: bool = False


@dataclass(frozen=True)
class ImportedStage:
    row: int
    seq: int
    code: str
    name: str
    phase: str
    zone_type: str
    visual_stage: str | None
    plan_start: date
    plan_end: date
    norm_duration_days: int
    work_codes: tuple[str, ...]
    predecessors: tuple[ImportedLink, ...]
    rule: dict[str, Any] | None
    basis: str


def read_table(filename: str, content: bytes) -> list[tuple[int, list[Any]]]:
    """Строки файла с номерами, как их видит человек в Excel или редакторе (с 1)."""
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "csv":
        return _read_csv(content)
    if extension == "xlsx":
        return _read_xlsx(content)
    raise PlanImportError([ImportIssue(None, None, "Поддерживаются файлы .csv и .xlsx")])


def _read_csv(content: bytes) -> list[tuple[int, list[Any]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel на русской Windows сохраняет CSV в cp1251.
        text = content.decode("cp1251")
    # Разделитель — по шапке: в ней нет данных, а в ячейках связей бывают и «;», и «,».
    header = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = max(";,\t", key=header.count)
    return list(enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1))


def _read_xlsx(content: bytes) -> list[tuple[int, list[Any]]]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except (InvalidFileException, zipfile.BadZipFile, KeyError, IndexError) as exc:
        raise PlanImportError([ImportIssue(None, None, "Файл XLSX не читается")]) from exc
    return [(n, list(row)) for n, row in enumerate(sheet.iter_rows(values_only=True), start=1)]


def parse_schedule(
    table: Sequence[tuple[int, Sequence[Any]]],
    templates: Mapping[str, TemplateStage],
    vocab: Vocabulary,
    calendar: WorkCalendar,
) -> list[ImportedStage]:
    """Этапы графика в порядке строк файла; при любой проблеме — PlanImportError со всеми."""
    header_row, columns = _header(table)
    issues: list[ImportIssue] = []
    stages: list[ImportedStage] = []
    for number, cells in table:
        if number <= header_row or all(_blank(c) for c in cells):
            continue
        values = {f: _cell(cells, i) for f, i in columns.items()}
        stage = _row(number, len(stages) + 1, values, templates, vocab, calendar, issues)
        if stage is not None:
            stages.append(stage)
    if not stages and not issues:
        issues.append(ImportIssue(None, None, "В файле нет ни одного этапа"))
    stages = _drop_missing_template_links(stages)
    issues += _check_codes_and_links(stages)
    if issues:
        raise PlanImportError(issues)
    return stages


def _header(table: Sequence[tuple[int, Sequence[Any]]]) -> tuple[int, dict[str, int]]:
    """Первая непустая строка — шапка; столбцы узнаются по названию, порядок любой."""
    for number, cells in table:
        if all(_blank(c) for c in cells):
            continue
        names = [str(c).strip().lower() if not _blank(c) else "" for c in cells]
        columns = {
            field: names.index(title)
            for field, titles in COLUMNS.items()
            for title in titles
            if title in names
        }
        missing = [COLUMNS[f][0] for f in REQUIRED if f not in columns]
        if missing:
            raise PlanImportError(
                [ImportIssue(number, None, f"В шапке нет столбцов: {', '.join(missing)}")]
            )
        return number, columns
    raise PlanImportError([ImportIssue(None, None, "Файл пуст")])


def _row(
    number: int,
    seq: int,
    values: Mapping[str, Any],
    templates: Mapping[str, TemplateStage],
    vocab: Vocabulary,
    calendar: WorkCalendar,
    issues: list[ImportIssue],
) -> ImportedStage | None:
    """Одна строка файла → этап; проблемы складываются в issues, этап тогда None."""
    found: list[ImportIssue] = []

    def problem(column: str, message: str) -> None:
        found.append(ImportIssue(number, column, message))

    code, name = _code(values.get("code")), _text(values.get("name"))
    if not code:
        problem("код", "Пустой код этапа")
    if not name:
        problem("наименование", "Пустое наименование этапа")
    start, end = _date(values.get("start")), _date(values.get("end"))
    if start is None:
        problem("начало", f"Дата не распознана: {values.get('start')!r}; ожидается ГГГГ-ММ-ДД")
    if end is None:
        problem("окончание", f"Дата не распознана: {values.get('end')!r}; ожидается ГГГГ-ММ-ДД")
    template = templates.get(code or "")

    phase = _text(values.get("phase")) or (template.phase if template else None)
    if phase is None:
        problem("фаза", "Фаза не указана, а шаблона этапа с таким кодом нет")
    elif phase not in vocab.phases:
        problem("фаза", f"Неизвестная фаза {phase!r}; допустимо: {list(vocab.phases)}")
    zone_type = _text(values.get("zone_type")) or (template.zone_type if template else None)
    if zone_type is None:
        problem("тип участка", "Тип участка не указан, а шаблона этапа с таким кодом нет")
    elif vocab.zone_roles.get(zone_type) != WORK_ROLE:
        problem("тип участка", f"На участке типа {zone_type!r} работы этапов не идут")
    visual = _text(values.get("visual_stage")) or (template.visual_stage if template else None)
    if visual is not None and visual not in vocab.stage_labels:
        problem("визуальная стадия", f"Неизвестная стадия {visual!r}")
    links = _links(values.get("predecessors"), template, problem)

    duration = 0
    if start and end:
        if end < start:
            problem("окончание", f"Окончание {end.isoformat()} раньше начала {start.isoformat()}")
        else:
            duration = count_working_days(calendar, start, end + timedelta(days=1))
            if duration == 0:
                problem("окончание", "В интервале этапа нет ни одного рабочего дня")
    if found:
        issues += found
        return None
    return ImportedStage(
        row=number,
        seq=seq,
        code=code,
        name=name,
        phase=phase,
        zone_type=zone_type,
        visual_stage=visual,
        plan_start=start,
        plan_end=end,
        norm_duration_days=duration,
        work_codes=template.work_codes if template else (code,),
        predecessors=links,
        rule=template.rule if template else None,
        basis=f"Импорт графика; правило и участок — шаблон этапа {code}"
        if template
        else "Импорт графика",
    )


def _links(raw: Any, template: TemplateStage | None, problem) -> tuple[ImportedLink, ...]:
    """Связи из столбца; пустая ячейка — связи шаблона, «-» — связей нет."""
    text = _text(raw)
    if text is None:
        links = template.predecessors if template else ()
        return tuple(ImportedLink(t.code, t.type, t.lag_days, from_template=True) for t in links)
    if text.lower() in NO_LINKS:
        return ()
    result = []
    for part in re.split(r"[;,]", text):
        match = LINK_PATTERN.match(part.strip())
        if not match:
            problem("связи", f"Связь {part.strip()!r} не разобрана; пример: «12.3.1 FS+2»")
            continue
        kind = (match["type"] or "FS").upper()
        result.append(ImportedLink(match["code"], kind, int(match["lag"] or 0)))
    return tuple(result)


def _drop_missing_template_links(stages: list[ImportedStage]) -> list[ImportedStage]:
    """Файл может содержать не все этапы шаблона: связи шаблона на отсутствующие не ставятся."""
    codes = {s.code for s in stages}
    return [
        replace(
            s,
            predecessors=tuple(p for p in s.predecessors if not p.from_template or p.code in codes),
        )
        for s in stages
    ]


def _check_codes_and_links(stages: Sequence[ImportedStage]) -> list[ImportIssue]:
    """Повторы кодов, связи на отсутствующие этапы, циклы — по всему графику."""
    issues: list[ImportIssue] = []
    seen: dict[str, int] = {}
    for stage in stages:
        if stage.code in seen:
            issues.append(
                ImportIssue(
                    stage.row, "код", f"Код {stage.code} уже был в строке {seen[stage.code]}"
                )
            )
        seen.setdefault(stage.code, stage.row)
    for stage in stages:
        for link in stage.predecessors:
            if link.code not in seen:
                issues.append(
                    ImportIssue(
                        stage.row, "связи", f"Связь на этап {link.code}, которого нет в файле"
                    )
                )
    if issues:
        return issues
    ids = {s.code: uuid5(NAMESPACE_URL, s.code) for s in stages}
    try:
        order_stages(
            [
                CpmStage(
                    ids[s.code],
                    s.plan_start,
                    s.plan_end,
                    tuple(Link(ids[p.code], p.type, p.lag_days) for p in s.predecessors),
                )
                for s in stages
            ]
        )
    except CpmError as exc:
        issues.append(ImportIssue(None, "связи", str(exc)))
    return issues


def _cell(cells: Sequence[Any], index: int) -> Any:
    return cells[index] if index < len(cells) else None


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any) -> str | None:
    if _blank(value):
        return None
    return str(value).strip()


def _code(value: Any) -> str | None:
    """Код этапа из ячейки: Excel превращает «10.11» в дату, а «12» — в число 12.0.

    Дата `d.m.<год>` восстанавливается в код `d.m` — то же правило, что для справочника работ
    (data/README.md). Код вида «12.10», ставший числом 12.1, восстановить нельзя: такой
    столбец в Excel нужно хранить как текст.
    """
    if isinstance(value, datetime | date):
        return f"{value.day}.{value.month}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date() if text else None
        except ValueError:
            continue
    return None
