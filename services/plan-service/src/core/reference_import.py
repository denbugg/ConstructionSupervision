"""Разбор справочника работ из XLSX заказчика: «Сводный перечень строительных работ ЛТЦ».

Устройство файла — data/README.md. Excel испортил часть кодов (даты вместо «10.1», число
вместо «12»), у 4-го уровня кодов нет вовсе: парсер восстанавливает коды и иерархию, а у
каждой строки запоминает номер строки файла — любое значение проверяется по первоисточнику.
Ошибки собираются по всему файлу сразу, как в импорте графика.
"""

import io
import re
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

# Отметка обязательности в файле заказчика — U+02C5, а не латинская «v». Пусто — «не
# обязательно»; отметка не запрещает работу для других типов объектов (ТЗ, п. 1).
MARK = "˅"
CODE_TITLE = "№ п/п"
NAME_TITLE = "Вид работ"
CODE_PATTERN = re.compile(r"^\d+(?:\.\d+)*$")
# Строка без кода — работа 4-го уровня внутри ближайшего кода сверху (data/README.md).
MAX_LEVEL = 4
# Разделитель кода, который парсер придумал сам: «10.2/3» не спутать с кодом из файла
# и с кодом, который заказчик может добавить позже («10.2.3»).
GENERATED_SEPARATOR = "/"


@dataclass(frozen=True)
class ReferenceIssue:
    row: int | None
    column: str | None
    message: str


class ReferenceImportError(ValueError):
    """Файл справочника не принят: список проблем с номерами строк."""

    def __init__(self, issues: Sequence[ReferenceIssue]) -> None:
        self.issues = list(issues)
        first = self.issues[0]
        where = f"строка {first.row}: " if first.row else ""
        super().__init__(f"{where}{first.message}")


@dataclass(frozen=True)
class ParsedWorkType:
    row: int
    code: str
    name: str
    level: int
    parent_code: str | None
    # Ключи — названия столбцов типов объектов из шапки файла.
    applicable: dict[str, bool]
    source: str
    # Код стоял в ячейке, а не сгенерирован для строки без кода.
    code_in_file: bool
    # Что Excel записал вместо кода, если код восстановлен из даты: для отчёта об импорте.
    restored_from: str | None = None


@dataclass(frozen=True)
class _Header:
    row: int
    code: int
    name: int
    # Номер столбца → название типа объекта.
    types: dict[int, str]


@dataclass
class _Cursor:
    """Состояние прохода по файлу: какие коды уже были и куда относить строки без кода."""

    seen: dict[str, int]
    # Файл и лист для work_type.source; номер строки дописывается к каждой строке.
    source: str
    last_code: str | None = None
    generated: int = 0


def read_sheet(content: bytes) -> tuple[str, list[tuple[int, list[Any]]]]:
    """Имя первого листа XLSX и его строки с номерами, как их видит человек в Excel (с 1)."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except (InvalidFileException, zipfile.BadZipFile, KeyError, IndexError, OSError) as exc:
        raise ReferenceImportError([ReferenceIssue(None, None, "Файл XLSX не читается")]) from exc
    rows = [(n, list(row)) for n, row in enumerate(sheet.iter_rows(values_only=True), start=1)]
    return sheet.title, rows


def parse_work_types(
    table: Sequence[tuple[int, Sequence[Any]]], *, filename: str, sheet: str
) -> list[ParsedWorkType]:
    """Виды работ в порядке строк файла; при любой проблеме — ReferenceImportError со всеми."""
    header = _header(table)
    cursor = _Cursor(seen={}, source=f"{filename}, лист «{sheet}», строка")
    issues: list[ReferenceIssue] = []
    parsed: list[ParsedWorkType] = []
    for number, cells in table:
        if number <= header.row or all(_blank(c) for c in cells):
            continue
        item, found = _row(number, cells, header, cursor)
        issues += found
        if item is not None and not found:
            parsed.append(item)
    if not parsed and not issues:
        issues.append(ReferenceIssue(None, None, "В файле нет ни одного вида работ"))
    if issues:
        raise ReferenceImportError(issues)
    return parsed


def code_sort_key(code: str) -> tuple[tuple[int, int], ...]:
    """Порядок справочника по коду: «10.2» < «10.2/1» < «10.2.1» < «10.10».

    Строковое сравнение поставило бы «10.10» перед «10.2». Строки без кода идут сразу за
    своим кодом, как в файле: там они стоят под ним.
    """
    base, _, generated = code.partition(GENERATED_SEPARATOR)
    key = tuple((1, int(part)) for part in base.split(".") if part.isdigit())
    return (*key, (0, int(generated))) if generated.isdigit() else key


def _header(table: Sequence[tuple[int, Sequence[Any]]]) -> _Header:
    """Шапка — строка со столбцом «Вид работ»; выше неё в файле заголовок справочника."""
    for number, cells in table:
        titles = [_text(c).lower() for c in cells]
        if NAME_TITLE.lower() not in titles:
            continue
        if CODE_TITLE.lower() not in titles:
            raise ReferenceImportError(
                [ReferenceIssue(number, None, f"В шапке нет столбца «{CODE_TITLE}»")]
            )
        code, name = titles.index(CODE_TITLE.lower()), titles.index(NAME_TITLE.lower())
        # Типы объектов — все остальные подписанные столбцы шапки: новый тип в файле
        # заказчика не требует правки кода.
        types = {i: _text(cells[i]) for i, t in enumerate(titles) if t and i not in (code, name)}
        if not types:
            raise ReferenceImportError(
                [ReferenceIssue(number, None, "В шапке нет столбцов типов объектов")]
            )
        names = list(types.values())
        repeated = sorted({t for t in names if names.count(t) > 1})
        if repeated:
            raise ReferenceImportError(
                [ReferenceIssue(number, None, f"Столбцы повторяются: {', '.join(repeated)}")]
            )
        return _Header(number, code, name, types)
    raise ReferenceImportError(
        [ReferenceIssue(None, None, "Не найдена шапка: строка со столбцом «Вид работ»")]
    )


def _row(
    number: int, cells: Sequence[Any], header: _Header, cursor: _Cursor
) -> tuple[ParsedWorkType | None, list[ReferenceIssue]]:
    """Одна строка файла → вид работ и проблемы строки."""
    found: list[ReferenceIssue] = []

    def problem(column: str, message: str) -> None:
        found.append(ReferenceIssue(number, column, message))

    raw = _cell(cells, header.code)
    code, restored_from = _code(raw, problem)
    name = _text(_cell(cells, header.name))
    if not name:
        problem(NAME_TITLE, "Пустое наименование работы")
    applicable = _marks(cells, header, problem)

    if code is not None:
        level, parent = _place(code, cursor.seen, problem)
        if code in cursor.seen:
            problem(CODE_TITLE, f"Код {code} уже был в строке {cursor.seen[code]}")
        # Код запоминается и при других ошибках строки: иначе её дети получат ложную
        # ошибку «нет родителя» поверх настоящей.
        cursor.seen.setdefault(code, number)
        cursor.last_code, cursor.generated, in_file = code, 0, True
    elif not _blank(raw):
        return None, found  # код не разобран, причина уже в found
    elif cursor.last_code is None:
        problem(CODE_TITLE, "Строка без кода выше первого кода: не к чему её отнести")
        return None, found
    else:
        cursor.generated += 1
        code = f"{cursor.last_code}{GENERATED_SEPARATOR}{cursor.generated}"
        level, parent, in_file = MAX_LEVEL, cursor.last_code, False
    item = ParsedWorkType(
        row=number,
        code=code,
        name=name,
        level=level,
        parent_code=parent,
        applicable=applicable,
        source=f"{cursor.source} {number}",
        code_in_file=in_file,
        restored_from=restored_from,
    )
    return item, found


def _code(value: Any, problem) -> tuple[str | None, str | None]:
    """Код из ячейки и то, из чего он восстановлен (только для дат).

    Excel прочитал «10.1» как 10 января: дата `d.m.<год>` → код `d.m`. Число «12» → «12».
    Дробное число восстановить нельзя: 12.1 могло быть и «12.1», и «12.10».
    """
    if _blank(value):
        return None, None
    if isinstance(value, datetime | date):
        return f"{value.day}.{value.month}", value.strftime("%d.%m.%Y")
    if isinstance(value, int | float) and not isinstance(value, bool):
        if float(value).is_integer():
            return str(int(value)), None
        problem(CODE_TITLE, f"Код стал числом {value}: неясно, «{value}» это или «{value}0»")
        return None, None
    text = str(value).strip().rstrip(".")
    if not CODE_PATTERN.match(text):
        problem(CODE_TITLE, f"Код {value!r} не разобран; ожидается вида «12.3.1.»")
        return None, None
    return text, None


def _place(code: str, seen: dict[str, int], problem) -> tuple[int, str | None]:
    """Уровень — число частей кода, родитель — код без последней части, строкой выше."""
    parts = code.split(".")
    if len(parts) > MAX_LEVEL:
        problem(CODE_TITLE, f"Код {code}: уровней больше {MAX_LEVEL}")
    parent = ".".join(parts[:-1]) or None
    if parent is not None and parent not in seen:
        problem(CODE_TITLE, f"Код {code}: выше нет строки с кодом родителя {parent}")
    return len(parts), parent


def _marks(cells: Sequence[Any], header: _Header, problem) -> dict[str, bool]:
    result = {}
    for index, title in header.types.items():
        value = _cell(cells, index)
        result[title] = not _blank(value) and str(value).strip() == MARK
        if not _blank(value) and not result[title]:
            problem(title, f"Отметка {value!r} не распознана: ожидается «{MARK}» или пусто")
    return result


def _cell(cells: Sequence[Any], index: int) -> Any:
    return cells[index] if index < len(cells) else None


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any) -> str:
    """Текст ячейки без переносов строк и двойных пробелов, которые оставил Excel."""
    return "" if _blank(value) else " ".join(str(value).split())
